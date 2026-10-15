"""Internal seminar deck (about 60 min, Korean): teleoperation survey -> LeRobot / ALOHA / ACT -> QuantGate -> AEP,
framed as the safety lifecycle of learned robot policies (the team's Safety Robotics focus for next year).
Every number comes from committed artifacts or from the cited papers' abstracts; every paper id was checked on alphaXiv.
Usage: <venv with python-pptx, PIL> python scripts/make_seminar.py <shots dir> <out.pptx>"""
import os, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

SHOTS, OUT = sys.argv[1:3]
FONT = "Malgun Gothic"
INK, INK2, MUTED, ACC, ACC2, WARN, BAD, PAPER, BAND = (RGBColor(20, 24, 28), RGBColor(60, 70, 78), RGBColor(110, 118, 124), RGBColor(0, 128, 140),
                                                       RGBColor(120, 170, 40), RGBColor(170, 120, 20), RGBColor(160, 50, 45), RGBColor(250, 251, 250), RGBColor(232, 240, 240))
prs = Presentation(); prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
PX = lambda v: Emu(int(v * 12192000 / 1920)); BL = prs.slide_layouts[6]; N = [0]

def tb(s, x, y, w, h, lines, size=20, color=INK, bold=False, align=None, anchor=None):
    box = s.shapes.add_textbox(PX(x), PX(y), PX(w), PX(h)); tf = box.text_frame; tf.word_wrap = True
    if anchor: tf.vertical_anchor = anchor
    for i, ln in enumerate(lines if isinstance(lines, list) else [lines]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        lvl, txt = (1, ln[2:]) if isinstance(ln, str) and ln.startswith("  ") else (0, ln)
        if isinstance(txt, tuple): txt, c, b = txt
        else: c, b = color, bold
        r = p.add_run(); r.text = txt; r.font.size = Pt(size - 3 * lvl); r.font.color.rgb = c; r.font.bold = b; r.font.name = FONT
        p.level = lvl; p.space_after = Pt(6)
        if align: p.alignment = align
    return box

def slide(title, sub=None, notes=""):
    s = prs.slides.add_slide(BL); s.background.fill.solid(); s.background.fill.fore_color.rgb = PAPER; N[0] += 1
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PX(14), prs.slide_height); bar.fill.solid(); bar.fill.fore_color.rgb = ACC; bar.line.fill.background()
    tb(s, 70, 40, 1700, 80, title, 32, INK, True)
    if sub: tb(s, 72, 112, 1760, 50, sub, 17, MUTED)
    tb(s, 1760, 1020, 120, 40, str(N[0]), 12, MUTED, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes; return s

def section(no, title, sub, notes=""):
    s = prs.slides.add_slide(BL); s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor(16, 24, 28); N[0] += 1
    tb(s, 160, 380, 1600, 80, f"PART {no}", 22, RGBColor(120, 200, 205), True); tb(s, 160, 440, 1600, 120, title, 44, RGBColor(240, 244, 244), True)
    tb(s, 160, 560, 1600, 80, sub, 20, RGBColor(170, 184, 186)); s.notes_slide.notes_text_frame.text = notes; return s

def box(s, x, y, w, h, text, fill=BAND, color=INK, size=16, bold=False, line=None):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PX(x), PX(y), PX(w), PX(h)); sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line: sh.line.color.rgb = line; sh.line.width = Pt(1.5)
    else: sh.line.fill.background()
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = ln; r.font.size = Pt(size if i == 0 else size - 3); r.font.bold = bold and i == 0; r.font.color.rgb = color; r.font.name = FONT
    return sh

def arrow(s, x1, y1, x2, y2):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, PX(x1), PX(y1), PX(x2), PX(y2)); c.line.color.rgb = MUTED; c.line.width = Pt(2)
    c.line._get_or_add_ln().append(__import__("lxml").etree.fromstring('<a:tailEnd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" type="triangle"/>'))

def pic(s, path, x, y, w, h):
    im = Image.open(path); r = min(w / im.width, h / im.height); ww, hh = im.width * r, im.height * r
    return s.shapes.add_picture(path, PX(x + (w - ww) / 2), PX(y + (h - hh) / 2), PX(ww), PX(hh))

def table(s, x, y, w, rows, widths, size=13, head_fill=ACC):
    t = s.shapes.add_table(len(rows), len(rows[0]), PX(x), PX(y), PX(w), PX(40 * len(rows))).table
    for j, cw in enumerate(widths): t.columns[j].width = PX(cw)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.text = ""; p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = str(val); r.font.size = Pt(size); r.font.name = FONT
            r.font.bold = i == 0; r.font.color.rgb = RGBColor(255, 255, 255) if i == 0 else INK
            c.fill.solid(); c.fill.fore_color.rgb = head_fill if i == 0 else (RGBColor(255, 255, 255) if i % 2 else RGBColor(240, 245, 245))
    return t

R = "artifacts/reports/figures"; G = "docs/media/gallery"; ST = "dashboard/setup"

# 1 title
s = prs.slides.add_slide(BL); s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor(16, 24, 28); N[0] += 1
tb(s, 140, 300, 1640, 200, "안전하게 배포할 수 있는 로봇 학습", 50, RGBColor(240, 244, 244), True)
tb(s, 140, 420, 1640, 120, "텔레오퍼레이션에서 배포 검증까지 — SO-101 실험(QuantGate)과 AEP", 26, RGBColor(170, 220, 222))
tb(s, 140, 560, 1640, 120, ["내부 세미나 · 약 60분", "Safety Robotics 관점: 데이터 → 학습·압축 → 배포 전 검증 → 운영 중 감시 → 사후 분석"], 18, RGBColor(170, 184, 186))
pic(s, "docs/assets/quantgate_logo.png", 1460, 760, 360, 220)
s.notes_slide.notes_text_frame.text = "오늘의 큰 질문: 학습된 로봇 정책을 언제 믿고 배포할 수 있는가. 텔레오퍼레이션 서베이로 시작해 우리가 SO-101로 한 실험과 AEP 논문까지, 내년 Safety Robotics 과제와 연결해서 봅니다."

# 2 agenda / lifecycle
s = slide("오늘의 질문: 학습된 정책을 언제 믿고 배포할 수 있는가?", "학습 정책의 안전 수명주기와 오늘 다룰 내용",
          "학습 정책은 규칙 기반 제어처럼 보장을 주지 않는다. 그래서 수명주기 각 단계에서 '무엇으로 믿을 근거를 만드는가'가 안전의 핵심이다. 오늘 이야기는 이 다섯 단계를 따라간다.")
stages = [("1. 데이터", "텔레오퍼레이션\n시연 품질·개입"), ("2. 학습·압축", "ACT 학습\n양자화 (AEP)"), ("3. 배포 전 검증", "블라인드 실기\n배포 전 점검 (QuantGate)"),
          ("4. 운영 중 감시", "런타임 모니터\n안전 필터"), ("5. 사후 분석", "에피소드 조사\n(LabLens)")]
for i, (a, b) in enumerate(stages):
    box(s, 80 + i * 360, 260, 320, 200, a + "\n" + b, fill=BAND if i != 2 else RGBColor(214, 236, 238), size=20, bold=True)
    if i < 4: arrow(s, 80 + i * 360 + 320, 360, 80 + (i + 1) * 360, 360)
table(s, 80, 540, 1760, [["시간", "파트", "내용"], ["0–5분", "도입", "왜 학습 정책의 안전이 어려운가"], ["5–20분", "PART 1", "텔레오퍼레이션: ALOHA·ACT, 인터페이스 서베이, 안전 관점(데이터 품질·개입)"],
                         ["20–28분", "PART 2", "LeRobot 구조와 우리 SO-101 셋업, 데이터 수집에서 배운 것"], ["28–42분", "PART 3", "QuantGate: 양자화된 정책의 배포 검증 실험과 결과"],
                         ["42–50분", "PART 4", "AEP: 고정 상태 정확도 ≠ 폐루프 성능"], ["50–60분", "PART 5", "운영 중 감시 연구 지도, 내년 Safety Robotics 방향, 질의"]], [200, 220, 1340], 14)

# 3 why hard
s = slide("학습 정책의 안전은 왜 어려운가", None,
          "핵심은 '성공률만으로는 안전을 말할 수 없다'는 것. SafeVLA-Bench는 성공해도 궤적 중 안전 위반이 숨는 '성공-안전 격차'를 지적한다. 우리 실험에서도 같은 작업, 같은 정책인데 그리퍼 마찰이나 시작 위치 3 cm만으로 성공률이 크게 바뀌었다.")
tb(s, 80, 180, 1000, 800, ["보장이 없다", "  신경망 정책은 규칙처럼 경계 조건을 증명할 수 없다",
                            "성공 ≠ 안전", "  목표는 달성해도 궤적 중 충돌·위험 접촉이 숨을 수 있다 (SafeVLA-Bench, 2606.00773)",
                            "분포 이동에 약하다", "  우리 SO-101: 시작 위치 3 cm 이동만으로 0/3, 그리퍼 패드 제거만으로 8/24",
                            "압축·배포가 동작을 바꾼다", "  양자화·컴파일러 설정이 같은 모델의 폐루프 행동을 바꾼다 (AEP)",
                            "평가가 비싸다", "  '실패 0'을 보이려면 수백 회 실기가 필요 — 25회로는 ±28%p 불확실"], 21)
pic(s, f"{G}/03_position_finetune.png", 1110, 200, 740, 500)
tb(s, 1110, 710, 740, 120, "같은 정책·같은 작업이 위치 이동에서 무너지고, 데이터로 고친 사례 (PART 3)", 14, MUTED)

# PART 1
section(1, "텔레오퍼레이션", "데이터를 만드는 수단이자, 사람이 개입하는 안전 장치",
        "텔레오퍼레이션은 두 가지 역할을 한다: 학습 데이터 수집, 그리고 배포 중 사람의 개입·보정. Safety 관점에서는 두 번째 역할이 점점 중요해지고 있다.")
s = slide("ALOHA와 ACT (Zhao et al., 2023)", "Learning Fine-Grained Bimanual Manipulation with Low-Cost Hardware · arXiv 2304.13705",
          "ALOHA는 저가 리더-팔로워 양팔 시스템. 사람이 리더 팔을 움직이면 팔로워가 따라 하고, 그때의 관측과 관절값이 시연 데이터가 된다. ACT는 한 번에 여러 스텝(chunk)의 행동을 예측해 누적 오차와 떨림을 줄인다. 우리 SO-101도 같은 리더-팔로워 구조, 같은 ACT를 썼다.")
tb(s, 80, 180, 820, 760, ["ALOHA: 저가 양팔 리더-팔로워 텔레오퍼레이션", "  사람은 리더 팔을 직접 움직이고, 팔로워가 같은 관절각을 따라간다",
                            "  카메라 영상 + 관절 상태 + 명령이 그대로 시연 데이터", "ACT (Action Chunking with Transformers)",
                            "  한 번에 k 스텝의 행동 묶음(chunk)을 예측 → 누적 오차 감소", "  CVAE로 사람 시연의 다양성(스타일)을 잠재변수로 흡수",
                            "  temporal ensembling으로 chunk 경계를 부드럽게", "우리 실험의 기본값", "  ACT 51.6M, chunk 100, 30 Hz, 카메라 2대(오버헤드·손목)"], 19)
box(s, 960, 220, 260, 110, "카메라 이미지\n(오버헤드·손목)", size=16); box(s, 960, 380, 260, 110, "관절 상태\n(6 DoF)", size=16)
box(s, 1300, 300, 260, 130, "ResNet 백본\n+ Transformer\n(CVAE 디코더)", fill=RGBColor(214, 236, 238), size=16, bold=True)
box(s, 1640, 300, 220, 130, "행동 chunk\nk = 100 스텝", size=16); arrow(s, 1220, 275, 1300, 365); arrow(s, 1220, 435, 1300, 365); arrow(s, 1560, 365, 1640, 365)
box(s, 960, 560, 900, 150, "실행: 30 Hz로 chunk를 소비하며, 남은 스텝이 적을 때 다음 chunk를 미리 요청\n우리 러너: 10 스텝 남으면 요청, 마지막 4 스텝은 이전·새 chunk를 섞어 전환 → 마감 200 ms", fill=BAND, size=15)
s = slide("ALOHA 계열의 확장", "Mobile ALOHA (2401.02117) · ALOHA 2 (2405.02292) · ALOHA Unleashed (2410.13126)",
          "Mobile ALOHA는 이동 베이스를 붙여 전신 텔레오퍼레이션으로 확장했고, 정적 ALOHA 데이터와 함께 학습(co-training)하면 적은 시연으로도 성능이 오른다는 점을 보였다. ALOHA 2는 하드웨어 내구성과 사용성을 개선해 데이터 수집 규모를 키웠고, ALOHA Unleashed는 대규모 시연과 diffusion 정책으로 정교한 작업까지 밀어붙였다.")
table(s, 80, 200, 1760, [["시스템", "무엇을 바꿨나", "의미"], ["Mobile ALOHA (2024)", "이동 베이스 + 전신 텔레오퍼레이션, 정적 데이터와 co-training", "이동 조작 시연을 저비용으로 수집"],
                         ["ALOHA 2 (2024)", "하드웨어 내구성·인체공학·카메라 개선", "대규모·장시간 수집의 신뢰성"], ["ALOHA Unleashed (2024)", "대규모 시연 + diffusion 정책", "모방학습만으로 정교한 양손 작업"],
                         ["GELLO (2023, 2309.13037)", "기존 산업용 팔에 맞춘 저가 리더 장치", "같은 퍼펫 방식을 다른 로봇에"], ["SO-100/101 + LeRobot", "수십만 원대 리더-팔로워, 오픈소스 스택", "우리가 쓴 구성"]], [420, 760, 580], 15)
tb(s, 80, 660, 1760, 300, ["Safety 관점 포인트", "  퍼펫 방식은 신체 차이(embodiment gap)가 거의 없어 시연이 그대로 실행 가능 → 데이터의 물리적 타당성이 높다",
                            "  대신 힘·접촉 피드백이 없거나 약해, 접촉이 많은 작업에서는 '위험한 시연'이 섞일 수 있다"], 19)
s = slide("텔레오퍼레이션 인터페이스 지도 (서베이)", "비용·신체 차이·피드백·확장성의 절충",
          "인터페이스는 크게 다섯 갈래. 리더-팔로워는 신체 차이가 없고, 비전·VR은 장비가 가볍지만 리타게팅이 필요하다. 로봇 없이 사람 시연을 모으는 UMI 계열은 규모에 강하지만 변환이 필요하다. 2026년에는 휴머노이드 전신과 촉각·햅틱이 가장 활발하다.")
table(s, 60, 180, 1800, [["갈래", "대표 연구 (arXiv)", "장점", "한계"],
                         ["A. 리더-팔로워", "ALOHA 2304.13705 · GELLO 2309.13037 · ModPack 2607.19479", "신체 차이 거의 없음, 저비용", "로봇마다 리더 필요, 힘 피드백 약함"],
                         ["B. 비전·VR 손 추적", "AnyTeleop 2307.04577 · Open-TeleVision 2407.01512 · Bunny-VisionPro 2407.03162 · ACE 2408.11805 · Phone2Act 2605.01948 · COBALT 2605.19138", "장비 가벼움, 원격·크라우드 가능", "리타게팅 오차, 지연"],
                         ["C. 로봇 없는 사람 시연", "UMI 2402.10329 · FastUMI 2409.19499 · DexCap 2403.07788 · DexUMI 2505.21864 · HoMMI 2603.03243 · HuMI 2602.06643", "규모·다양성, 현장 수집", "로봇으로 옮기는 변환 필요"],
                         ["D. 휴머노이드 전신", "H2O 2403.04436 · OmniH2O 2406.08858 · HumanPlus 2406.10454 · Teleopit 2608.01834 · SPOT 2609.07933", "전신 조작·이동", "균형·안전 제약이 큼"],
                         ["E. 정교한 손·촉각·햅틱", "Human-level Dexterous Teleop 2607.11481 · DITTO 2609.19196 · XRoboToolKit-T 2609.16437", "접촉 품질, 힘 감각", "장비 복잡·고가"]], [300, 900, 300, 300], 12)
s = slide("안전 관점에서 본 텔레오퍼레이션", "데이터 품질 · 사람의 개입 · 접촉 피드백",
          "Safety Robotics에서 텔레오퍼레이션은 '데이터 수집기'를 넘어 '안전 장치'가 된다. 수집 단계에서 나쁜 시연을 거르고, 배포 중에는 사람이 개입해 보정하며, 그 개입이 다시 학습 데이터가 되는 순환(DAgger 계열)이 핵심이다.")
for i, (h, body) in enumerate((("수집 단계: 데이터 품질", "Closing the Loop in Teleoperation (Siemens, 2605.26349)\n에피소드 단위로 품질을 평가해 수집 중에 피드백\n→ 나쁜 시연이 정책의 위험 행동이 되기 전에 차단"),
                               ("배포 단계: 개입과 보정", "DITTO-X (2610.00781): 순방향·역방향 텔레오퍼레이션으로 배포 중 개입\nReward-DAgger (2610.04054): 진행도 보상 모델이 '언제 사람에게 넘길지' 판단\n→ 개입이 곧 on-policy 보정 데이터"),
                               ("접촉 단계: 힘·촉각", "DITTO (2609.19196), XRoboToolKit-T (2609.16437)\n촉각·햅틱으로 접촉을 느끼며 시연 → 과도한 힘·위험 접촉이 데이터에 덜 섞임"))):
    box(s, 80 + i * 600, 200, 560, 420, h + "\n" + body, fill=BAND, size=19, bold=True)
tb(s, 80, 680, 1760, 300, ["우리 경험과의 연결", "  빈 에피소드 녹화(q를 누른 채 시작), 리더 팔 전압 오류, 모터 패킷 오류 → 수집 품질 문제는 실제로 자주 일어난다",
                            "  다음 단계 아이디어: 배포한 정책의 실제 rollout을 FP32 교사로 라벨링하는 on-policy 감사·증류 (DITTO-X·DAgger 계열과 같은 방향)"], 18)

# PART 2
section(2, "LeRobot과 우리 셋업", "오픈소스 스택으로 SO-101 데이터 수집부터 배포까지", "")
s = slide("LeRobot 구조", "LeRobot: An Open-Source Library for End-to-End Robot Learning (2602.22818) · Robot Learning: A Tutorial (2510.12403)",
          "LeRobot은 하드웨어 추상화(로봇·텔레오퍼레이터·카메라), 표준 데이터셋, 정책 구현, 명령줄 도구를 한 묶음으로 제공한다. 우리는 record로 시연을 모으고, ACT를 Nebius에서 학습한 뒤, 정책을 TensorRT로 바꿔 Jetson에서 실행했다.")
for i, (h, b) in enumerate((("하드웨어", "robots · teleoperators · cameras\nSO-101 follower/leader, OpenMV 플러그인"), ("데이터", "LeRobotDataset v3\nparquet + mp4 + meta"),
                            ("정책", "ACT · Diffusion · π0 · SmolVLA …\n정규화 통계 포함 체크포인트"), ("도구", "lerobot-record · teleoperate\nreplay · train · eval"))):
    box(s, 80 + i * 450, 200, 410, 260, h + "\n" + b, size=18, bold=True, fill=BAND)
tb(s, 80, 500, 1760, 400, ["우리 사용 흐름", "  lerobot-record (리더로 시연, 데모 120개) → LOCK-1 분할 80/20/20 → Nebius Jobs에서 ACT 학습 (L40S, 1시간 57분)",
                            "  → ONNX → TensorRT FP32/FP16/INT8 엔진 26개 → Jetson Orin Nano Super에서 30 Hz 실행",
                            "실무 팁", "  카메라 플러그인(OpenMV)처럼 하드웨어는 플러그인으로 붙인다",
                            "  캘리브레이션 파일이 로봇마다 다르다 → 장비를 옮기면 같이 옮겨야 한다"], 19)
s = slide("LeRobotDataset v3 — 분석할 때의 함정", "여러 에피소드가 한 파일에 들어간다", "")
tb(s, 80, 180, 900, 800, ["파일 구조", "  data/chunk-000/file-000.parquet: 여러 에피소드의 행", "  videos/<cam>/chunk-000/file-000.mp4: 여러 에피소드의 프레임",
                           "  meta/episodes: 에피소드별 행 범위와 영상 구간(from/to timestamp)", "  meta/info.json: features, fps, 코덱, 경로 규칙",
                           "함정", "  파일명으로 에피소드를 짝지으면 안 된다 → 메타데이터의 offset을 읽어야 한다",
                           "  영상은 AV1 (재생기에 따라 미지원)", "  단위: 관절 도(degree), 그리퍼 0–100 위치값 (힘이 아님)"], 19)
table(s, 1020, 200, 840, [["필드 (에피소드 1)", "값"], ["data 파일", "chunk-000 / file-000"], ["전역 index", "702 – 1499 (798 프레임)"], ["영상 구간", "23.4 s – 50.0 s"], ["분할", "train"]], [420, 420], 15)
box(s, 1020, 520, 840, 300, "정책 rollout 기록은 다르다\n프레임 i = 제어 스텝 i (정확)\n그러나 영상은 30 fps로 저장되고, 실제 스텝 간격은\n중앙값 33 ms · 최대 약 200 ms → 시간은 실제 시각 열로", fill=RGBColor(250, 236, 214), size=17, bold=True)
s = slide("우리 셋업: SO-101 + Jetson Orin Nano Super", "과제: 268 g 콜라캔을 시작 표시에서 집어 컵받침 위에 세운다", "")
pic(s, f"{ST}/setup_photo.jpg", 80, 180, 560, 780)
pic(s, f"{ST}/task_start_top.jpg", 700, 200, 560, 420); pic(s, f"{ST}/task_end_top.jpg", 1300, 200, 560, 420)
tb(s, 700, 640, 1160, 340, ["리더-팔로워 SO-101 (Feetech STS3215), 오버헤드 C922 640×480 + OpenMV 손목 320×200",
                             "정책은 Jetson에서 TensorRT로 30 Hz 실행, 학습·감사는 Nebius 클라우드", "데이터: 시연 120개 + 위치 데모 50개 (HF: chans/quantgate-so101-coaster)"], 18)
s = slide("데이터 수집과 하드웨어에서 배운 것", "안전 사고는 대부분 '평범한 문제'에서 시작된다",
          "실제 운영에서 문제는 모델보다 하드웨어·운영 실수가 먼저였다. 그리퍼 마찰 하나가 성공률을 23/24에서 8/24로 바꿨다. 이런 조건은 반드시 배포 검토에 '검증된 조건'으로 적어야 한다.")
table(s, 80, 190, 1760, [["사건", "원인", "안전 관점 교훈"], ["빈 에피소드가 녹화됨", "q를 누른 채 시작", "수집 직후 에피소드 품질 자동 점검 필요"],
                         ["리더 팔 전압 오류, 모터 패킷 오류", "전원·케이블", "하드웨어 상태 감시 없이는 데이터·실험 모두 오염"],
                         ["미끄럼 패드 제거 → 8/24", "그리퍼 마찰", "정책이 아니라 '배포 조건'이 성공률을 좌우"],
                         ["시작 위치 3 cm → 0/3", "학습 분포 밖", "검증된 조건의 범위를 명시해야 함"],
                         ["엔진 빌드마다 지연이 다름", "TensorRT 비결정성", "타이밍 캐시 고정 — 같은 파일이 같은 행동을 보장해야 함"]], [520, 400, 840], 15)

# PART 3
section(3, "QuantGate: 배포 전 검증", "양자화된 정책을 배포해도 되는가 — 증거로 답하기",
        "QuantGate는 Nebius×NVIDIA 해커톤 출품작. '이 엔진을 배포해도 되나?'에 대해 비용·실기 결과·증거의 강도를 나란히 보여주는 검토 도구다.")
s = slide("문제와 접근", "후보 26개: FP32 · FP16 · INT8 보정 변형 24개 (같은 체크포인트)", "")
for i, (h, b) in enumerate((("1. 대상 장치 비용", "Jetson에서 직접 지연·에너지 측정"), ("2. 오프라인 감사", "보류 상태 2,000개에서 점수\n(실기 전 순위 고정)"),
                            ("3. 블라인드 실기", "봉인·인터리브·운영자 블라인드\n결과는 해시 고정 후 해제"), ("4. 배포 검토", "추천 ≠ 승인\n근거가 부족하면 '아직'"))):
    box(s, 80 + i * 450, 220, 410, 230, h + "\n" + b, size=20, bold=True, fill=BAND if i != 2 else RGBColor(214, 236, 238))
    if i < 3: arrow(s, 80 + i * 450 + 410, 335, 80 + (i + 1) * 450, 335)
tb(s, 80, 520, 1760, 420, ["연구 무결성 장치 (안전 검증에 그대로 필요한 것)", "  사전등록: 판정 규칙을 결과보다 먼저 커밋", "  블라인딩: 운영자는 어느 엔진인지 모름, 화면에 지연도 숨김",
                            "  봉인과 해시: 실험 기록을 열기 전에 해시를 고정 → 결과를 보고 규칙을 바꿀 수 없음",
                            "  판정: GO / NO-GO / INSUFFICIENT_EVIDENCE (Newcombe 구간, 허용 오차 10%p)"], 19)
s = slide("결과 1 — 비용은 대상 장치에서 재야 한다", "Jetson Orin Nano Super, 15 W, TensorRT 10.3 (trtexec)", "")
pic(s, f"{SHOTS}/cost.png", 80, 180, 1100, 760)
tb(s, 1220, 220, 640, 700, ["FP16: 11.6 ms, 82 mJ", "INT8: 16.8 ms, 101 mJ", "FP32: 53.4 ms, 403 mJ", ("→ 이 장치에서는 INT8이 FP16보다 느리고 에너지도 더 씀", ACC, True),
                             "이름(정밀도)이 비용을 말해주지 않는다", "같은 모델을 Hailo-8로 옮기면 보정 크기 효과가 반대로 나왔다 (오프라인 탐색)"], 19)
s = slide("결과 2 — '실패가 안 보였다' ≠ 안전", "오프라인 점수는 실기를 예측하지 못했다", "")
tb(s, 80, 190, 900, 760, ["감사가 고른 INT8 셀 vs 무작위 셀: 73/75 vs 73/75", "  오프라인 순위와 실제 배치 오차 상관: Spearman 0.02–0.08",
                           "FP32 24/25, INT8 10개 셀 194/200 — 차이 없어 보임", ("  그러나 25회로는 차이 구간이 ±28%p → 판정은 '근거 부족'", BAD, True),
                           "  10%p 안을 보이려면 팔당 약 110–120회", "안전 메시지", "  검증의 결론은 '안전하다'가 아니라 '어디까지 확인됐다'여야 한다"], 19)
pic(s, f"{R}/fig_e1_forest.png", 1000, 190, 860, 760)
s = slide("결과 3 — 위험은 배포 조건에 있었고, 데이터로 고쳤다", "위치 파인튜닝: 블라인드·인터리브, 위치당 12 대 12", "")
pic(s, f"{R}/fig_ft_position.png", 80, 180, 900, 520); pic(s, f"{SHOTS}/t1_pair_frame.png", 1000, 200, 860, 480)
tb(s, 80, 720, 1760, 260, ["−4…+4 cm 시연 50개로 20k 스텝 파인튜닝 (Nebius Jobs, 25분) → 학습에 없던 위치에서 평가",
                            "+3 cm 12/12 vs 3/12 (+75%p), −3 cm 10/12 vs 2/12 (+67%p): 사전등록 기준 개선 확인 · 원래 위치는 10/12 vs 12/12로 아직 근거 부족",
                            "한계: 두 정책의 동작이 눈에 보여 운영자 블라인딩은 부분적 — 판정은 자동 판정과 영상 재판정"], 17)
s = slide("배포 전 점검 — 소프트웨어 인터록", "깨진 빌드는 실기 전에 막는다", "")
pic(s, f"{SHOTS}/preflight.png", 80, 180, 1100, 780)
tb(s, 1220, 200, 640, 760, ["순서대로 네 가지 (명시적 GO만 통과)", "  엔진 해시 = 기록된 빌드", "  정규화 파일 매니페스트 (러너가 직접 거부)", "  실시간 마감 200 ms", "  코드가 정한 판정",
                             "결과", "  정상 FP16: HOLD (시험 더 필요)", "  FT 정책: HOLD (개선 확인 ≠ 승인)", "  파일 바꿔치기·매니페스트 누락·엔진 교체: BLOCK", "  Hailo 분할 실행 p95 264 ms: BLOCK", "  검은 화면 보정 빌드: BLOCK"], 18)
s = slide("AI 감독자·판정기도 '검증 대상 부품'이다", "Cosmos-Reason2 (Nebius AI Endpoint) · Nemotron 3 Super (Token Factory)",
          "AI를 안전 판단에 쓰려면 그 AI부터 검증해야 한다. 첫 설정의 Cosmos는 실패를 거의 못 잡았다. 녹화로 프롬프트를 고르고 따로 검증한 뒤, 새 실기에서 다시 확인했다. 보고서도 숫자와 판정은 코드가 쓰고 모델은 말로만 설명하게 했다.")
table(s, 80, 190, 1100, [["Cosmos-Reason2 끝 상태 판정기", "실패 감지", "오경보"], ["첫 설정 (한 단어 답)", "1/21", "0/252"], ["검증된 프롬프트, 보류 집합", "20/21", "5/252 (노란 캔 4/33)"],
                         ["같은 프롬프트, 새 실기 72회", "21/23", "1/49"], ["2 Hz 감독자 (CONTINUE/ABORT)", "실패 감지 못함", "—"]], [560, 260, 280], 15)
tb(s, 80, 470, 1100, 500, ["교훈", "  문구에 민감 → 배포 전 그 문구 그대로 보류 데이터로 재검증", "  같은 결론: FailBench (2609.03611) — 최고 VLM도 균형 정확도 0.77, 모호하면 '성공'으로 기운다",
                            "  모델 출력은 숫자·판정을 쓰지 못하게 하고 코드가 검증 (보고서 v3, 질의응답)"], 18)
pic(s, f"{R}/fig_frames.png", 1220, 190, 640, 760)

# PART 4
section(4, "AEP 논문", "Better Fixed-State Fidelity Need Not Improve Closed-Loop Performance in Quantized Neural Policies",
        "AEP는 QuantGate의 이론적 배경. 고정된 상태에서 FP32와 더 비슷하게 행동한다고 해서, 실제 폐루프에서 더 잘하는 것은 아니라는 주장이다.")
s = slide("AEP: 고정 상태 정확도 ≠ 폐루프 성능", "양자화된 정책에서 '더 비슷한 행동'이 '더 좋은 결과'를 보장하지 않는다",
          "숫자는 초안 초록 기준. Hailo FetchReach 계열에서 고정 상태 평균 행동 오차가 16% 더 작은 아티팩트가 FP32 대비 반환 손실은 34배 컸다. AdaRound를 켠 컴파일러 개입은 8개 규칙 모두 RMSE를 줄였지만 평균 반환은 5개에서 감소했다. 그래서 '측정된 실행 파일을 실제 실행 경로에서' 검증하는 절차를 제안한다.")
tb(s, 80, 190, 1000, 780, ["질문", "  FP32와 행동이 더 비슷하면 양자화 후 폐루프 성능도 더 좋은가?",
                            "설계", "  학습된 정책은 고정, 보정 데이터·컴파일러 설정만 바꿔 Edge TPU · Hailo에서 실행 (시뮬레이션 제어 과제)",
                            "핵심 결과 (초안 초록)", "  고정 상태 행동 오차가 16% 작은 아티팩트가 반환 손실은 34배 큼",
                            "  AdaRound 개입: 8개 규칙 모두 RMSE 감소, 그러나 평균 반환은 5개에서 감소",
                            "제안", "  정확도 기반 '후보 선택'과 과제 허용 오차 기반 '승인'을 분리", "  측정한 실행 파일을 실제 실행 경로에서 평가"], 19)
box(s, 1120, 200, 740, 340, "물리 로봇 근거 (QuantGate → AEP)\nSO-101 실기 G6, 사전등록 v1\n오프라인 감사 최상위 vs 무작위: 73/75 vs 73/75\n오프라인 순위 vs 배치 오차: Spearman ≈ 0", fill=RGBColor(214, 236, 238), size=18, bold=True)
box(s, 1120, 580, 740, 300, "Safety 관점의 의미\n'오프라인 지표가 좋다'는 안전 인증이 아니다\n배포 경로 그대로의 실행 검증이 필요하다", fill=BAND, size=18, bold=True)

# PART 5
section(5, "운영 중 감시와 내년 방향", "Safety Robotics: 배포 이후에도 근거를 만든다", "")
s = slide("운영 중 감시 연구 지도 (2025–2026)", "실패 감지 · 안전 필터 · 안전 벤치마크 · 판정기 검증", "")
table(s, 60, 180, 1800, [["갈래", "대표 연구 (arXiv)", "무엇을 하나"],
                         ["실행 중 실패 감지", "SAFE 2506.09937 · VLA-FAIL 2606.21386 · FailureSpot 2609.04277 · Foresight 2606.23085 · FARM 2609.11445 · Rewind-IL 2604.16683", "정책 내부 신호·세계모델 잠재로 실패를 미리·시점 단위로 경고"],
                         ["접촉 전 감시", "ContactGuard 2608.13438 · PATCH 2606.16690", "접촉 전에 위험을 예측 (손목 카메라)"],
                         ["안전 필터", "Attention-Guided Safety Filter 2606.09749 · WBAG 2610.01083 · VLPSA 2609.22462 · Latent Safety Filters 2608.00315", "정책 출력을 충돌 회피 제약으로 수정"],
                         ["안전 벤치마크", "SafeVLA-Bench 2606.00773 · SafeStage 2609.21223 · ForesightSafety-VLA 2606.27079", "성공률 뒤에 숨은 안전 위반 측정"],
                         ["판정기·검증기", "FailBench 2609.03611 · No Free Checker (검증기 서베이) 2609.09250", "VLM 성공 판정의 신뢰도와 한계"],
                         ["배포 신뢰성", "Deployment-Time Reliability of Learned Robot Policies 2603.11400 · OGAM 2610.05878", "체계적 시험과 런타임 보증의 연결"]], [330, 960, 510], 13)
s = slide("우리 작업과 내년 Safety Robotics의 연결", "수명주기 단계별: 한 일 → 남은 공백 → 다음 과제", "")
table(s, 60, 180, 1800, [["단계", "우리가 한 일", "남은 공백", "내년 과제 후보"],
                         ["데이터", "SO-101 시연 120+50, 품질 사고 기록", "수집 중 자동 품질 점검 없음", "에피소드 품질 점검 + LabLens식 조사 도구"],
                         ["학습·압축", "AEP: 정확도 ≠ 폐루프, Hailo 이식성 탐색", "압축이 바꾸는 '동작 품질' 미측정", "INT8 그리퍼 떨림(약 2배) 같은 행동 지표의 사전등록 검증"],
                         ["배포 전 검증", "QuantGate: 블라인드 실기, 구간 판정, 배포 전 점검", "승인까지 필요한 시험 규모가 큼", "on-policy 감사로 실기 비용 줄이기"],
                         ["운영 중 감시", "Cosmos 감독자 그림자 실행 (실패 감지 못함)", "감독자 프롬프트 미검증", "검증된 런타임 모니터 + 안전 필터, 개입 → 보정 데이터 (DAgger)"],
                         ["사후 분석", "운영자 노트, 영상 재판정, 실패 유형 라벨", "사건 단위 근거 연결 부족", "영상·궤적·가설을 잇는 조사 도구 (LabLens)"]], [260, 520, 460, 560], 14)
tb(s, 60, 760, 1800, 200, ["표준과의 연결: 산업용 로봇 안전 ISO 10218, 협동 로봇 ISO/TS 15066 — 학습 정책을 이 틀에 넣으려면 '검증된 조건'과 '근거의 강도'를 문서로 남기는 절차가 먼저 필요"], 16, MUTED)
s = slide("정리: 다섯 가지 메시지", None, "")
tb(s, 100, 190, 1720, 800, [("1. 텔레오퍼레이션은 데이터 수집기이자 안전 장치다", ACC, True), "  수집 품질, 배포 중 개입, 접촉 피드백이 모두 안전과 직결",
                             ("2. 비용과 동작은 대상 장치에서 재야 한다", ACC, True), "  Jetson에서는 INT8이 FP16보다 느렸다",
                             ("3. '실패가 안 보였다'는 안전 증거가 아니다", ACC, True), "  25회로는 ±28%p — 결론은 '어디까지 확인됐다'로",
                             ("4. 위험은 모델보다 배포 조건에 있었다", ACC, True), "  그리퍼 마찰·시작 위치. 찾아서 데이터로 고치고 사전등록 규칙으로 재시험",
                             ("5. AI 감시자도 검증해야 하는 부품이다", ACC, True), "  첫 설정 1/21 → 검증 후 21/23. 문구 하나에 결과가 바뀐다"], 21)
s = slide("참고문헌 (텔레오퍼레이션)", None, "")
tb(s, 80, 160, 1760, 860, ["ALOHA/ACT 2304.13705 · Mobile ALOHA 2401.02117 · ALOHA 2 2405.02292 · ALOHA Unleashed 2410.13126 · GELLO 2309.13037 · ModPack 2607.19479",
                            "AnyTeleop 2307.04577 · Open-TeleVision 2407.01512 · Bunny-VisionPro 2407.03162 · ACE 2408.11805 · Phone2Act 2605.01948 · COBALT 2605.19138",
                            "UMI 2402.10329 · FastUMI 2409.19499 · DexCap 2403.07788 · DexUMI 2505.21864 · HoMMI 2603.03243 · Humanoid Manipulation Interface 2602.06643",
                            "H2O 2403.04436 · OmniH2O 2406.08858 · HumanPlus 2406.10454 · Teleopit 2608.01834 · SPOT 2609.07933",
                            "Towards Human-level Dexterous Teleoperation 2607.11481 · DITTO 2609.19196 · DITTO-X 2610.00781 · XRoboToolKit-T 2609.16437",
                            "Closing the Loop in Teleoperation (데이터 품질) 2605.26349 · Reward-DAgger 2610.04054 · TeleOpBench 2505.12748",
                            "Teleoperation of Humanoid Robots: A Survey 2301.04317 · LeRobot 2602.22818 · Robot Learning: A Tutorial 2510.12403"], 15)
s = slide("참고문헌 (안전·검증)", None, "")
tb(s, 80, 160, 1760, 860, ["실패 감지: SAFE 2506.09937 · VLA-FAIL 2606.21386 · FailureSpot 2609.04277 · Foresight 2606.23085 · FARM 2609.11445 · Rewind-IL 2604.16683 · ContactGuard 2608.13438 · PATCH 2606.16690",
                            "안전 필터: Attention-Guided Safety Filter 2606.09749 · WBAG 2610.01083 · VLPSA 2609.22462 · Latent Safety Filters 2608.00315",
                            "벤치마크: SafeVLA-Bench 2606.00773 · SafeStage 2609.21223 · ForesightSafety-VLA 2606.27079 · FailBench 2609.03611 · No Free Checker 2609.09250",
                            "배포·압축: Deployment-Time Reliability 2603.11400 · OGAM 2610.05878 · PreDE 2609.19441 · VLAQuantBench 2609.25376 · FoldQuantVLA 2609.24433 · 8 GB Bimanual ACT 2608.03938",
                            "우리 작업: QuantGate (github.com/goodboychan/nebius-physical-ai, 10/26 공개 예정) · AEP 초안 (비공개)"], 15)
prs.save(OUT); print("saved", OUT, N[0], "slides")
