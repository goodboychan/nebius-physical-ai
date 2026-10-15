"""Scene plan for re-editing the submission video in PowerPoint (animated graphs, transitions).
Writes, under <out>:
  data/*.csv                  the numbers behind every graph or card on screen, from committed artifacts
  quantgate_scene_plan.pptx   one slide per video segment: time, length, a still from the current video (v6), the
                              on-screen chip, the narration spoken in that span, the media source, and the data file
Usage: <venv with av, PIL, python-pptx> python scripts/make_scene_plan.py <video v6.mp4> <out dir>"""
import csv, io, json, math, os, re, sys
import av
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from g6_analysis import wilson, newcombe

VID, OUT = sys.argv[1:3]; D = f"{OUT}/data"; os.makedirs(D, exist_ok=True)
R = "artifacts/reports"; W = lambda k, n: [round(x * 100, 1) for x in wilson(k, n)]

NOTES = []
def wcsv(name, head, rows, note):   # header on the first line (PowerPoint/Excel chart import); sources go to data/README.md
    with open(f"{D}/{name}", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(head); w.writerows(rows)
    NOTES.append(f"- `{name}`: {note}")

# 1. engine cost on the Jetson (trtexec, 15 W)
cost = [("VC-000", "FP32", 53.4, 403, "power_bench"), ("VC-025", "FP16", 11.6, 82, "power_bench"), ("FT-FP16", "FP16, fine-tuned", 11.6, 82, "power_bench_ft"),
        ("VC-002", "INT8, calib 16, per-channel", 16.8, 101, "power_bench"), ("VC-005", "INT8, calib 16, per-tensor", 17.1, 101, "power_bench"),
        ("VC-020", "INT8, calib 1024, per-channel", 16.8, 101, "power_bench")]
wcsv("01_engine_cost.csv", ["engine", "precision", "gpu_ms_median", "energy_mJ_per_inference", "source"], cost,
     "Jetson Orin Nano Super, 15 W, TensorRT 10.3, trtexec back-to-back; artifacts/reports/power_bench*/SUMMARY.md. Only these engines were benched.")
# 2. G6 real-arm success by cell
A = json.load(open(f"{R}/g6_analysis.json"))
rows = sorted(([c, v["k"], v["n"], round(v["rate"] * 100, 1), *[round(x * 100, 1) for x in v["wilson95"]], v["arms"][0]] for c, v in A["cells"].items()), key=lambda r: r[0])
wcsv("02_g6_success_by_cell.csv", ["cell", "successes", "trials", "success_pct", "wilson95_lo_pct", "wilson95_hi_pct", "arm"], rows,
     "G6 main study, blinded and interleaved, 245 trials; artifacts/reports/g6_analysis.json")
# 3. cost vs evidence (FP16 run and G6 cells with a measured cost)
T = json.load(open(f"{R}/d3_v3_verdicts.json")); rows = []
for r in T["rows"]:
    if r.get("gpu_ms_trtexec") and r.get("success") and r["context"].startswith(("G6", "FP16")):
        k, n = r["success"]; rows.append([r["candidate"], r["context"], r["gpu_ms_trtexec"], k, n, round(k / n * 100, 1), *W(k, n), r["verdict"]])
g6ref = next(r for r in T["rows"] if r["context"].startswith("G6"))["reference"]; rows.append(["VC-000", "G6 FP32 reference", 53.4, g6ref[0], g6ref[1], round(g6ref[0] / g6ref[1] * 100, 1), *W(*g6ref), "reference"])
wcsv("03_cost_vs_success.csv", ["candidate", "context", "gpu_ms", "successes", "trials", "success_pct", "wilson95_lo_pct", "wilson95_hi_pct", "verdict"], rows,
     "fig_cost_evidence.png: on-device GPU time vs real-arm success; verdicts from d3_v3_verdicts.json")
# 4. audit-picked vs random INT8 cells
e = A["EP-1"]
wcsv("04_audit_vs_random.csv", ["group", "cells", "successes", "trials", "success_pct"],
     [["audit top picks", " ".join(e["group1"]), e["k1"], e["n1"], round(e["k1"] / e["n1"] * 100, 1)], ["random picks", " ".join(e["group2"]), e["k2"], e["n2"], round(e["k2"] / e["n2"] * 100, 1)]],
     f"EP-1, G6; difference Newcombe 95% [{e['newcombe95'][0] * 100:+.1f}, {e['newcombe95'][1] * 100:+.1f}] pp")
# 5. position fine-tune
F, F3 = json.load(open(f"{R}/ft_analysis.json")), json.load(open(f"{R}/ft3_analysis.json")); rows = []
for blk, pos in (("T2", "-3 cm"), ("T1", "+3 cm")):
    p = F[blk]["primary"]; rows.append([pos, blk, *p["ft"], *W(*p["ft"]), *p["old"], *W(*p["old"]), round(p["diff_pp"], 1), *[round(x, 1) for x in p["newcombe95_pp"]], 95, "improvement detected"])
q = F["T0"]["secondary"]; d = q["E1"]
rows.insert(1, ["0 cm (training mark)", "T0", *q["test_k_n"], *W(*q["test_k_n"]), *q["ref_k_n"], *W(*q["ref_k_n"]), round(-d["degradation_pp"], 1), round(-d["ci_pp"][1], 1), round(-d["ci_pp"][0], 1), 97.5, "insufficient evidence"])
q = F3["pooled_T0_T0R"]; d = q["E1"]
rows.insert(2, ["0 cm (training mark)", "T0 + T0R pooled", *q["test_k_n"], *W(*q["test_k_n"]), *q["ref_k_n"], *W(*q["ref_k_n"]), round(-d["degradation_pp"], 1), round(-d["ci_pp"][1], 1), round(-d["ci_pp"][0], 1), 97.5, "insufficient evidence"])
wcsv("05_position_finetune.csv", ["start", "block", "ft_successes", "ft_trials", "ft_wilson95_lo_pct", "ft_wilson95_hi_pct", "old_successes", "old_trials",
     "old_wilson95_lo_pct", "old_wilson95_hi_pct", "ft_minus_old_pp", "ci_lo_pp", "ci_hi_pp", "ci_level_pct", "reading"], rows,
     "ft_analysis.json, ft3_analysis.json; +-3 cm: improvement rule (lower bound > 0, 95%); training mark: non-inferiority at 10 pp (97.5%)")
# 6. AI judge
L = json.load(open(f"{R}/ft_cosmos_live_judge.json"))
wcsv("06_ai_judge.csv", ["setup", "failures_caught", "failures", "false_alarms", "successes"],
     [["first setup (one-word answer), held-out recordings", 1, 21, 0, 252], ["validated prompt C1 p1, held-out recordings", 20, 21, 5, 252],
      ["same prompt, live new fine-tune trials", L["failure_recall"][0], L["failure_recall"][1], L["false_alarm"][0], L["false_alarm"][1]]],
     "Cosmos-Reason2 end-of-trial judge; artifacts/cosmos_judge/eval_scores.jsonl and artifacts/reports/ft_cosmos_live_judge.json")

open(f"{D}/README.md", "w").write("# Graph data for the video (from committed artifacts)\n\n" + "\n".join(NOTES) + "\n")
# scene plan
src = open("scripts/make_video.py").read(); REL = eval(src[src.index("REL = [") + 6: src.index("]\nSEG")] + "]", {"C": "C", "E": "E", "R": "R", "SHOTS": "SHOTS", "T1L": "T1L", "LOGO": "video/assets/quantgate_logo_intro.mp4"})
ST = json.load(open("docs/narration/narration_v3_timing.json"))
srt = open("docs/narration/narration_v3.srt").read().strip().split("\n\n"); ts = lambda s: sum(float(x) * m for x, m in zip(s.replace(",", ".").split(":"), (3600, 60, 1)))
SUB = [(ts(b.split("\n")[1].split(" --> ")[0]), ts(b.split("\n")[1].split(" --> ")[1]), " ".join(b.split("\n")[2:])) for b in srt]
DATA = {"cost.png": "01_engine_cost.csv", "success.png": "02_g6_success_by_cell.csv", "fig_cost_evidence.png": "03_cost_vs_success.csv", "contrasts.png": "04_audit_vs_random.csv",
        "fig_ft_position.png": "05_position_finetune.csv", "Cosmos": "06_ai_judge.csv", "pos_strip.png": "05_position_finetune.csv (demo positions −4…+4 cm)"}
TIPS = {"image": "그래프 → 데이터 CSV로 차트 재작성 후 애니메이션(막대·구간 순차 등장)", "clip": "영상 클립: 원본 파일에서 같은 구간 사용", "card": "텍스트 카드: 줄 단위 순차 등장",
        "intro": "로고 소개 (썸네일·로고 의미)", "logo": "마무리 로고"}
cont = av.open(VID); vs = cont.streams.video[0]
def still(t):
    cont.seek(int(t / vs.time_base), stream=vs)
    for f in cont.decode(vs):
        if float(f.pts * vs.time_base) >= t - 0.02: b = io.BytesIO(); f.to_image().convert("RGB").resize((640, 360)).save(b, "JPEG", quality=85); b.seek(0); return b
prs = Presentation(); prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000); PX = lambda v: Emu(int(v * 12192000 / 1920))
def tb(s, x, y, w, h, text, size, bold=False, color=RGBColor(20, 24, 28)):
    t = s.shapes.add_textbox(PX(x), PX(y), PX(w), PX(h)).text_frame; t.word_wrap = True
    for i, ln in enumerate(text.split("\n")):
        p = t.paragraphs[0] if i == 0 else t.add_paragraph(); r = p.add_run(); r.text = ln; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = "Malgun Gothic"
s = prs.slides.add_slide(prs.slide_layouts[6]); tb(s, 80, 80, 1760, 100, "QuantGate 영상 씬 정리 (v6 기준)", 36, True)
tb(s, 80, 200, 1760, 700, f"총 {len(REL)}개 구간 · 길이 {cont.duration / 1e6:.1f}초 · 내레이션 narration_v3 (am_michael)\n\n각 장: 시간, 화면 정지 컷, 화면 칩 문구, 그 구간 내레이션, 원본 미디어, 그래프 데이터 CSV\n"
   "그래프 데이터: data/01–06 *.csv (출처는 data/README.md)\n\n전환 효과·애니메이션을 넣어도 각 장면의 시작 시각과 길이를 지키면 내레이션이 그대로 맞습니다.\n길이를 바꾸면 이 표의 시작 시각 대신 장면 순서로 음성을 다시 배치합니다.", 20)
for n, (k, a, b, kind, args, chipt) in enumerate(REL, 1):
    t0, t1 = ST[k] + a, ST[k] + b; s = prs.slides.add_slide(prs.slide_layouts[6])
    tb(s, 60, 30, 1800, 60, f"{n:02d} · 씬 {k} · {t0:6.1f}–{t1:6.1f}초 ({t1 - t0:.1f}초) · {kind}", 26, True)
    s.shapes.add_picture(still((t0 + t1) / 2), PX(60), PX(120), PX(960), PX(540))
    media = args[0] if args and isinstance(args[0], str) else (args[0].split("\n")[0] if args else "—")
    data = next((v for kk, v in DATA.items() if kk in str(args)), "—")
    narr = " ".join(x[2] for x in SUB if x[1] > t0 + 0.05 and x[0] < t1 - 0.05)
    tb(s, 1060, 120, 800, 120, "화면 칩: " + (chipt or "—"), 16, True)
    tb(s, 1060, 260, 800, 160, "원본: " + str(media).replace("SHOTS/", "대시보드 캡처 ").replace("R/", "artifacts/reports/figures/").replace("C/", "dashboard/clips/").replace("E/", "video/edit/"), 14)
    tb(s, 1060, 430, 800, 80, "데이터: " + data, 14, True, RGBColor(0, 128, 140))
    tb(s, 1060, 520, 800, 80, "제안: " + TIPS.get(kind, ""), 14, False, RGBColor(110, 118, 124))
    tb(s, 60, 690, 1800, 360, "내레이션: " + (narr or "—"), 16)
    s.notes_slide.notes_text_frame.text = narr
prs.save(f"{OUT}/quantgate_scene_plan.pptx"); print("saved", OUT, len(REL), "segments")
