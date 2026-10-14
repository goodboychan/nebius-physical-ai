"""Build an editable PowerPoint version of the submission video: one slide per picture segment of scripts/make_video.py.
Each slide has its media (picture, or a pre-cut H.264 clip), the on-screen chip as a text box, the segment's narration
sentences in the speaker notes, and an auto-advance time equal to the segment length, so a slideshow follows the
narration track. Shapes are separate objects so the user can add and time animations in PowerPoint.
The narration audio is not embedded: scene WAVs and the full track are delivered next to the deck.
Usage: <venv with av, PIL, python-pptx> python scripts/make_deck.py <shots dir> <out dir>"""
import os, sys
import av
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from lxml import etree
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import phone_clips as pc

SHOTS, OUT = sys.argv[1:3]; os.makedirs(f"{OUT}/media", exist_ok=True)
src = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "make_video.py")).read()
g = {"__name__": "timeline", "__file__": os.path.join(os.path.dirname(os.path.abspath(__file__)), "make_video.py")}
exec(compile(src[:src.index("def fit(")], "make_video_timeline", "exec"), g)   # constants and the segment list only
SEG, SUBS = g["SEG"], g["SUBS"]
W, H = Emu(12192000), Emu(6858000)   # 16:9
PX = lambda v: Emu(int(v * 12192000 / 1920))
DARK, INK, ACC = RGBColor(16, 20, 19), RGBColor(230, 234, 233), RGBColor(18, 160, 168)

prs = Presentation(); prs.slide_width, prs.slide_height = W, H; blank = prs.slide_layouts[6]

def text(slide, x, y, w, h, s, size=20, color=INK, bold=False, mono=False):
    tb = slide.shapes.add_textbox(PX(x), PX(y), PX(w), PX(h)); tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(s.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
        if mono: r.font.name = "Consolas"
    return tb

def picture(slide, path, box=(80, 110, 1760, 860), crop=None):
    im = Image.open(path)
    if crop: im = im.crop((0, crop[0], im.width, min(im.height, crop[1]))); path = f"{OUT}/media/{os.path.basename(path)[:-4]}_crop.png"; im.save(path)
    r = min(box[2] / im.width, box[3] / im.height); w, h = im.width * r, im.height * r
    return slide.shapes.add_picture(path, PX(box[0] + (box[2] - w) / 2), PX(box[1] + (box[3] - h) / 2), PX(w), PX(h))

def cut(i, path, s0, speed, rot, dur):
    out = f"{OUT}/media/seg{i:02d}.mp4"; poster = f"{OUT}/media/seg{i:02d}_poster.png"
    o = av.open(out, "w", options={"movflags": "+faststart"}); s = o.add_stream("libx264", rate=30); s.width, s.height, s.pix_fmt = 1280, 720, "yuv420p"
    s.options = {"crf": "24", "preset": "medium"}; first = None
    for _, im in pc.frames(path, s0, dur * speed, fps=30 / speed):
        if rot: im = im.rotate(rot)
        im = im.convert("RGB"); r = min(1280 / im.width, 720 / im.height); fr = Image.new("RGB", (1280, 720), (16, 20, 19))
        im2 = im.resize((int(im.width * r), int(im.height * r))); fr.paste(im2, ((1280 - im2.width) // 2, (720 - im2.height) // 2))
        if first is None: first = fr; fr.save(poster)
        for p in s.encode(av.VideoFrame.from_image(fr)): o.mux(p)
    for p in s.encode(): o.mux(p)
    o.close(); return out, poster

def advance(slide, seconds):
    ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    tr = etree.SubElement(slide._element, f"{{{ns}}}transition"); tr.set("advTm", str(int(seconds * 1000)))
    # p:transition must come right after p:clrMapOvr (or p:cSld)
    slide._element.remove(tr); cs = slide._element.find(f"{{{ns}}}clrMapOvr")
    (cs if cs is not None else slide._element.find(f"{{{ns}}}cSld")).addnext(tr)

for i, (t0, t1, kind, args, chip) in enumerate(SEG):
    sl = prs.slides.add_slide(blank); bg = sl.background.fill; bg.solid(); bg.fore_color.rgb = DARK
    if kind == "image": picture(sl, args[0], crop=args[1:] if len(args) == 3 else None)
    elif kind == "clip":
        path, s0, speed, rot = args; mp4, poster = cut(i, path, s0, speed, rot, t1 - t0)
        box = (80, 110, 1760, 860); r = min(box[2] / 1280, box[3] / 720); w, h = 1280 * r, 720 * r
        sl.shapes.add_movie(mp4, PX(box[0] + (box[2] - w) / 2), PX(box[1] + (box[3] - h) / 2), PX(w), PX(h), poster_frame_image=poster, mime_type="video/mp4")
    elif kind == "card":
        mono = "\n\n" in args[0]; text(sl, 160, 300, 1600, 500, args[0], size=22 if mono else 40, bold=not mono, mono=mono)
    elif kind == "intro":
        picture(sl, "docs/assets/quantgate_thumbnail.png", box=(0, 0, 1920, 1080))
        sc = 1080 / 1024; ox = (1920 - 1536 * sc) / 2; P = lambda x, y: (ox + x * sc, y * sc)   # thumbnail -> slide px
        from pptx.enum.shapes import MSO_SHAPE
        for txt, box, at in (("Q · QuantGate and quantization", (600, 255, 932, 567), (120, 150)),
                             ("the gap at the corner: the gate", (848, 455, 928, 555), (1190, 380)),
                             ("a candidate leaving the gate:\ncleared for its next test, not approved", (902, 507, 976, 580), (1190, 520))):
            (x0, y0), (x1, y1) = P(box[0], box[1]), P(box[2], box[3])
            rr = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PX(x0), PX(y0), PX(x1 - x0), PX(y1 - y0)); rr.fill.background()
            rr.line.color.rgb = RGBColor(255, 255, 255); rr.line.width = Pt(2.5)
            tx, ty = P(*at); text(sl, tx, ty, 460, 140, txt, size=18)   # separate shapes: animate each with its sentence
    else:
        picture(sl, "docs/assets/quantgate_logo.png", box=(510, 200, 900, 420))
        from pptx.enum.text import PP_ALIGN
        for y, t_, sz, col, b in ((700, "Evidence before deployment.", 34, INK, True), (800, "github.com/goodboychan/nebius-physical-ai", 22, RGBColor(174, 184, 183), False)):
            tb = text(sl, 260, y, 1400, 80, t_, size=sz, color=col, bold=b); tb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    if chip: text(sl, 40, 24, 1840, 60, chip, size=18, color=ACC, bold=True)
    said = [s for s in SUBS if s[0] < t1 and s[1] > t0]
    sl.notes_slide.notes_text_frame.text = (f"[{t0:.1f}–{t1:.1f} s, {t1 - t0:.1f} s] {kind}\n" + "\n".join(f"{a:.1f}s  {x}" for a, b, x in said))
    advance(sl, t1 - t0)
    print(f"slide {i + 1:2d}  {t0:6.1f}-{t1:6.1f}  {kind}", flush=True)
prs.save(f"{OUT}/quantgate_video_deck.pptx"); print("saved", f"{OUT}/quantgate_video_deck.pptx")
